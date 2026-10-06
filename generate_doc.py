import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Title
title = doc.add_heading('西南山地山洪地灾与水保遥感智能监测系统\n—— 核心功能与业务逻辑说明书', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph('本系统（“山地慧眼”）是一套面向西南复杂山区，集“空天地采集 - 双域检测 - 图斑解析 - 风险评估 - 核查反馈 - 样本沉淀”于一体的全流程业务原型系统。系统构建了从 AI 自动发现隐患到人工现场核查处置，再到数据反哺 AI 的完整闭环。')

# Section 1
doc.add_heading('一、 双域智能检测 (AI Remote Sensing Detection)', level=1)
doc.add_paragraph('系统的核心 AI 计算引擎，负责处理遥感影像并自动提取地表变化图斑。', style='Intense Quote')
p1 = doc.add_paragraph()
p1.add_run('1. 时序双影像对比：').bold = True
p1.add_run('支持 T1（历史）与 T2（当前）高分辨率遥感影像的同步输入与可视化对比。')
p2 = doc.add_paragraph()
p2.add_run('2. 深度学习双引擎驱动：').bold = True
p2.add_run('内嵌 TIDAL-Net（侧重地灾/滑坡/地表扰动提取）与 FSD-Net（侧重水土保持要素/水体提取）双模型。')
p3 = doc.add_paragraph()
p3.add_run('3. 实时推理与图斑提取：').bold = True
p3.add_run('基于 FastAPI 后端，实现全自动的图像预处理、切片预测、形态学后处理，并在界面上实时绘制变化区域的边界框（Bounding Box），计算图斑面积与中心坐标。')
p4 = doc.add_paragraph()
p4.add_run('4. 数据状态无缝流转：').bold = True
p4.add_run('检测完毕后，结果（任务流水号、图斑编号、变化类型）会自动缓存并无缝流转至下游的“风险预估”节点，消除人工二次录入。')

# Section 2
doc.add_heading('二、 综合风险预估 (Risk Evaluation)', level=1)
doc.add_paragraph('系统的研判大脑，耦合地质环境因子与动态气象预报，进行防灾减灾级别的综合判定。', style='Intense Quote')
p5 = doc.add_paragraph()
p5.add_run('1. 一键模拟提取 GIS 环境因子：').bold = True
p5.add_run('支持一键基于图斑坐标获取静态地质特征。后台实现了动态生成逼真的高程（m）、坡度（°）、地形起伏、NDVI（植被覆盖指数）及距道路/断层距离等关键数据，模拟真实生产中的 GIS 切片服务。')
p6 = doc.add_paragraph()
p6.add_run('2. 动态降雨条件考量：').bold = True
p6.add_run('支持输入历史6日前期有效雨量、当前短时降雨强度、持续时间及未来24小时预报，全面贴合山洪灾害爆发的真实气象诱因。')
p7 = doc.add_paragraph()
p7.add_run('3. 专业级交叉判定算法：').bold = True
p7.add_run('后端引擎采用多权重结合算法（坡度30%、起伏20%、道路扰动15%等）算出“静态易发性”；同时结合降雨计算“动态红橙黄蓝预警”。两者交叉矩阵输出最终的“综合风险等级（极高/高/中/低）”。')
p8 = doc.add_paragraph()
p8.add_run('4. 高危自动派单机制：').bold = True
p8.add_run('对于判定为“高”或“极高”风险的图斑，系统会自动生成核查工单，静默派发至“核查任务”队列，给出“立即核查/优先核查”等行动建议。')

# Section 3
doc.add_heading('三、 核查任务工作台 (Inspection Tasks)', level=1)
doc.add_paragraph('系统的基层管理节点，用于呈现高风险预警并记录现场排查结果。', style='Intense Quote')
p9 = doc.add_paragraph()
p9.add_run('1. 态势大屏仪表盘：').bold = True
p9.add_run('顶部直观展示“待核查”、“高风险任务”、“已完成”核心数据指标，方便管理者一屏掌握当前排查进度。')
p10 = doc.add_paragraph()
p10.add_run('2. 多维智能过滤系统：').bold = True
p10.add_run('支持按任务编号、图斑编号模糊检索，并支持按风险等级、任务状态进行组合筛选。')
p11 = doc.add_paragraph()
p11.add_run('3. 全景上下文详情追溯：').bold = True
p11.add_run('侧滑抽屉设计，完整展示当时触发报警的各项因素（如：前期有效降雨偏高、坡度较大、动态降雨达到橙色预警等），为现场人员提供核查依据。')
p12 = doc.add_paragraph()
p12.add_run('4. 现场结果录入与一键导出：').bold = True
p12.add_run('允许核查人员在系统中录入现场结论（如误报、真变化）并归档。提供 Excel 一键导出台账功能。')

# Section 4
doc.add_heading('四、 反馈样本与审核管理 (Feedback & AI Iteration)', level=1)
doc.add_paragraph('系统的核心护城河：将业务核查数据沉淀为结构化的模型训练集，形成 AI 迭代飞轮。', style='Intense Quote')
p13 = doc.add_paragraph()
p13.add_run('1. 数据分布态势感知：').bold = True
p13.add_run('顶部统计面板实时呈现样本库内的“全部反馈”、“误报(负样本)”、“漏报(正样本)”、“确认正确”及“待审核”数据量，为模型迭代提供直观的数据储备参考。')
p14 = doc.add_paragraph()
p14.add_run('2. 极简式关联录入与影像补充：').bold = True
p14.add_run('告别手工填表，下拉选择关联的核查任务即可自动带入模型置信度、风险等级等所有前置上下文。支持现场实景照片与手工标注截图的同步上传。')
p15 = doc.add_paragraph()
p15.add_run('3. 三级数据防污审核机制：').bold = True
p15.add_run('新增反馈默认进入“待审核”状态。只有经由管理员点击“确认”后，才会流转为有效样本；同时支持“作废”操作，严格杜绝现场脏数据直接污染下游训练集。')
p16 = doc.add_paragraph()
p16.add_run('4. 优质训练集导出反哺：').bold = True
p16.add_run('支持将“已确认”的优质样本一键导出为标准化的 Excel 清单，供算法工程师脱机进行大模型的增量再训练，完美闭环！')

doc.add_paragraph('\n\n*本说明文档由系统自动生成，适用于大挑（挑战杯）项目申报、软件著作权申请及答辩演示说明。*').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.save('E:\\PycharmProjects\\website\\山地慧眼_功能说明书.docx')